#!/usr/bin/env python3
"""
Debug script to analyze the actual resume.docx structure.
"""

from docx import Document
import sys

print("=" * 70)
print("Resume Structure Analysis")
print("=" * 70)

try:
    doc = Document("input/resume.docx")

    print(f"\nTotal paragraphs: {len(doc.paragraphs)}\n")

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        # Get style info
        style_name = para.style.name if para.style else "None"

        # Check for numbering (bullet/list formatting)
        has_numbering = False
        if para._element.pPr is not None:
            numPr = para._element.pPr.numPr
            if numPr is not None:
                has_numbering = True

        # Show first character
        first_char = text[0] if text else ""
        first_char_code = ord(first_char) if first_char else 0

        print(f"Para {i}:")
        print(f"  Style: {style_name}")
        print(f"  Has numbering: {has_numbering}")
        print(f"  First char: '{first_char}' (code: {first_char_code})")
        print(f"  Text length: {len(text)}")
        print(f"  Text preview: {text[:80]}")
        print()

except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
