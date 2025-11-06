#!/usr/bin/env python3
"""
Create a test resume.docx from the provided resume text.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Set narrow margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

# Header - Name
name = doc.add_paragraph('SADSADSADAS')
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.runs[0].bold = True
name.runs[0].font.size = Pt(16)

# Contact info
contact = doc.add_paragraph('+1 (469) 487-XXXX | sxgggggggg@gmail.com | linkedin.com/in/asdasdsa | github.com/sadasdasd')
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.runs[0].font.size = Pt(10)

# Work Experience Section
doc.add_paragraph()
work_exp = doc.add_paragraph('WORK EXPERIENCE')
work_exp.runs[0].bold = True
work_exp.runs[0].font.size = Pt(12)

# Ghost Founder LLC
doc.add_paragraph()
company1_header = doc.add_paragraph('Ghost Founder LLC \t\t\t\t\t\t Dallas, TX')
company1_header.runs[0].bold = True

position1 = doc.add_paragraph('Software Engineer\t\t\t\t\t Aug 2023 – Present')
position1.runs[0].italic = True

# Bullets for Ghost Founder
bullets1 = [
    "Integrated AWS IoT and MQTT pipelines with secure token-based access, improving device data integrity and ensuring frictionless telemetry delivery across all deployed units",
    "Refactored frontend state management using Context API and React hooks, improving render efficiency by ~25% and enabling responsive control across 20+ deployed devices",
    "Designed reusable UI components for device health, usage analytics, and alert notifications, enhancing internal tool UX consistency and reducing front-end code duplication by 40%",
    "Partnered with hardware and backend teams to standardize API schemas and data contracts, improving integration reliability and simplifying multi-service debugging",
    "Containerized Python microservices using Docker, improving portability and CI/CD consistency, and aligning deployments across edge devices and cloud instances",
    "Secured $5 million in pre-seed funding by demonstrating system reliability, scalability, and operational efficiency",
    "Applied software QA methodologies by writing test plans and defining requirements, improving test coverage by 30% and catching 95% of defects pre-release"
]

for bullet in bullets1:
    p = doc.add_paragraph(bullet, style='List Bullet')
    p.runs[0].font.size = Pt(10)

# American Airlines
doc.add_paragraph()
company2_header = doc.add_paragraph('American Airlines\t\t\t\t\t\t Dallas, TX')
company2_header.runs[0].bold = True

position2 = doc.add_paragraph('Student Researcher\t\t\t\t Jan 2023 – May 2023')
position2.runs[0].italic = True

bullets2 = [
    "Deployed Python ETL pipelines on AWS to process 10M+ operational logs with secure handling of sensitive data, enabling compliant and efficient model-training workflows",
    "Applied machine learning models (Scikit-learn, Pandas, NumPy) to identify key drivers of customer satisfaction, uncover operational anomalies, and improve data interpretation for leadership reviews",
    "Developed customer churn prediction models using Random Forest and K-Nearest Neighbors, achieving 92% accuracy through hyperparameter tuning and optimization"
]

for bullet in bullets2:
    p = doc.add_paragraph(bullet, style='List Bullet')
    p.runs[0].font.size = Pt(10)

# Projects Section
doc.add_paragraph()
projects = doc.add_paragraph('PROJECTS')
projects.runs[0].bold = True
projects.runs[0].font.size = Pt(12)

# Cal Project
doc.add_paragraph()
project1_header = doc.add_paragraph('Cal | Personalized AI Fitness Coach for iOS')
project1_header.runs[0].bold = True

bullets3 = [
    "Deployed an iOS fitness app in Swift, integrating Apple HealthKit and REST APIs, and released a TestFlight beta that drove strong adoption and active usage among pilot users",
    "Integrated secure REST APIs and user auth flows with AWS Lambda and Core ML, ensuring privacy of health data while enabling real-time, adaptive workout recommendations",
    "Automated model retraining and logging with AWS S3 and CloudWatch, enhancing observability of ML performance and ensuring continuous, reliable recommendations"
]

for bullet in bullets3:
    p = doc.add_paragraph(bullet, style='List Bullet')
    p.runs[0].font.size = Pt(10)

# Bloom Project
doc.add_paragraph()
project2_header = doc.add_paragraph('Bloom | Depression Prediction using AI')
project2_header.runs[0].bold = True

bullets4 = [
    "Created an AI-powered depression prediction tool using Convolutional Neural Networks (CNNs) and multiple regression models, achieving 92% accuracy with only a 6% false-negative rate to minimize missed at-risk cases",
    "Deployed Python/React portal on AWS with Postgres backing for secure data handling and API auth, improving reliability and user privacy for clinical AI predictions"
]

for bullet in bullets4:
    p = doc.add_paragraph(bullet, style='List Bullet')
    p.runs[0].font.size = Pt(10)

# Moodify Project
doc.add_paragraph()
project3_header = doc.add_paragraph('Moodify | Playlist generation using API based on moods')
project3_header.runs[0].bold = True

bullets5 = [
    "Constructed a Node.js backend with mood classification models (~85% accuracy), automating mood tagging and enabling scalable real-time playlist generation",
    "Integrated AWS API Gateway and Lambda for secure mood-based playlist requests, improving latency by 35% and standardizing API observability across frontend and backend"
]

for bullet in bullets5:
    p = doc.add_paragraph(bullet, style='List Bullet')
    p.runs[0].font.size = Pt(10)

# Technical Skills
doc.add_paragraph()
skills = doc.add_paragraph('TECHNICAL SKILLS')
skills.runs[0].bold = True
skills.runs[0].font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph('Programming Languages: Python, SQL, Java, JavaScript, TypeScript, C, C#, HTML, CSS, PHP, R')
doc.add_paragraph('Frameworks & Libraries: React, Node.js, Express, Core ML, Axios, Scikit-learn, Pandas, NumPy, PyTorch')
doc.add_paragraph('Tools and Technologies: Git, Docker, MQTT, Grafana, MySQL, InfluxDB, AWS IoT, REST APIs, Tableau, Power BI, ETL')
doc.add_paragraph('Platforms & Systems: Linux, Raspberry Pi OS, TestFlight')
doc.add_paragraph('Certification: AWS Cloud Solution Architect - Associate')

# Education
doc.add_paragraph()
education = doc.add_paragraph('EDUCATION')
education.runs[0].bold = True
education.runs[0].font.size = Pt(12)

doc.add_paragraph()
school1_header = doc.add_paragraph('The University of Texas at Dallas\t\t\t\t Dallas, TX')
school1_header.runs[0].bold = True
doc.add_paragraph('Master of Science, Information Technology and Management; GPA: 3.6/4.0\t Aug 2021 – May 2023')

doc.add_paragraph()
school2_header = doc.add_paragraph('University of Mumbai\t\t\t\t\t Mumbai, India')
school2_header.runs[0].bold = True
doc.add_paragraph('Bachelor of Engineering, Computer Engineering\t\t\t Aug 2017 – May 2021')

# Save the document
doc.save('input/resume.docx')
print("✓ Resume created at: input/resume.docx")
print(f"✓ Total bullet points: {len(bullets1) + len(bullets2) + len(bullets3) + len(bullets4) + len(bullets5)}")
